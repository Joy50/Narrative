from collections import defaultdict
from datetime import timedelta
from io import BytesIO

from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from config.pagination import paginate
from .fetcher import has_active_fetch_job, start_fetch_job
from .models import Content, Country, FacebookPage, FetchJob, Keyword, Publisher
from .nlp import categorize_news, classify_sentiment, summarize
from .serializers import (
    ContentSerializer,
    CountrySerializer,
    FacebookPageSerializer,
    FetchJobSerializer,
    KeywordSerializer,
    PublisherSerializer,
)


class OverviewView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        today = timezone.now().date()
        countries = Country.objects.annotate(news_count=Count('articles'))
        return Response({
            'articles': Content.objects.count(),
            'today': Content.objects.filter(publication_date=today).count(),
            'publishers': Publisher.objects.count(),
            'keywords': Keyword.objects.filter(is_active=True).count(),
            'jobs': FetchJob.objects.count(),
            'countries': CountrySerializer(countries, many=True).data,
        })


class CountryListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        qs = Country.objects.annotate(news_count=Count('articles'))
        return paginate(qs, request, CountrySerializer)

    def post(self, request):
        serializer = CountrySerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class PublisherListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        publishers = Publisher.objects.select_related('country')
        if request.query_params.get('all') == 'true':
            return Response(PublisherSerializer(publishers, many=True).data)
        return paginate(publishers, request, PublisherSerializer)

    def post(self, request):
        serializer = PublisherSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class KeywordListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return paginate(Keyword.objects.all(), request, KeywordSerializer)

    def post(self, request):
        serializer = KeywordSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class KeywordDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):
        keyword = get_object_or_404(Keyword, pk=pk)
        serializer = KeywordSerializer(keyword, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    def delete(self, request, pk):
        get_object_or_404(Keyword, pk=pk).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class FacebookPageListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return paginate(FacebookPage.objects.select_related('country'), request, FacebookPageSerializer)

    def post(self, request):
        serializer = FacebookPageSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class FacebookPageDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):
        page = get_object_or_404(FacebookPage, pk=pk)
        serializer = FacebookPageSerializer(page, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    def delete(self, request, pk):
        get_object_or_404(FacebookPage, pk=pk).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class NewsListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        days_raw = (request.query_params.get('days') or 'all').strip().lower()
        qs = Content.objects.select_related('country', 'publisher')
        if days_raw not in ('all', '0', ''):
            days = max(int(days_raw), 1)
            qs = qs.filter(publication_date__gte=timezone.now().date() - timedelta(days=days))
        country_id = request.query_params.get('country')
        sentiment = request.query_params.get('sentiment')
        category = request.query_params.get('category')
        query = request.query_params.get('q')
        if country_id:
            qs = qs.filter(country_id=country_id)
        if sentiment:
            qs = qs.filter(sentiment=sentiment)
        if category:
            qs = qs.filter(category=category)
        if query:
            qs = qs.filter(
                Q(title__icontains=query) | Q(summary__icontains=query) | Q(body__icontains=query)
                | Q(publisher__name__icontains=query) | Q(country__name__icontains=query)
                | Q(matched_keyword__icontains=query)
            )
        return paginate(qs, request, ContentSerializer)

    def post(self, request):
        data = request.data
        publisher = get_object_or_404(Publisher, pk=data.get('publisher'))
        country = get_object_or_404(Country, pk=data.get('country') or publisher.country_id)
        text = data.get('body') or data.get('title') or ''
        article = Content.objects.create(
            title=data.get('title', '')[:500],
            publisher=publisher,
            body=data.get('body', ''),
            url=data.get('url'),
            publication_date=data.get('publication_date') or timezone.now().date(),
            country=country,
            summary=data.get('summary') or summarize(text),
            sentiment=data.get('sentiment') or classify_sentiment(text),
            category=data.get('category') or categorize_news(text),
            matched_keyword=data.get('matched_keyword', ''),
            source_kind='manual',
        )
        return Response(ContentSerializer(article).data, status=status.HTTP_201_CREATED)


class NewsDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        return Response(ContentSerializer(get_object_or_404(Content, pk=pk)).data)


class StatisticsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, country_id=None):
        days = max(int(request.query_params.get('period') or 7), 1)
        end = timezone.now().date()
        start = end - timedelta(days=days - 1)
        qs = Content.objects.filter(publication_date__gte=start)
        country = None
        if country_id:
            country = get_object_or_404(Country, pk=country_id)
            qs = qs.filter(country=country)
        by_country = {row['country__name']: row['count'] for row in qs.values('country__name').annotate(count=Count('id'))}
        sentiments = {'Positive': 0, 'Negative': 0, 'Neutral': 0}
        for row in qs.values('sentiment').annotate(count=Count('id')):
            if row['sentiment'] in sentiments:
                sentiments[row['sentiment']] = row['count']
        comparison = []
        for item in Country.objects.all():
            scoped = Content.objects.filter(country=item, publication_date__gte=start)
            total = scoped.count()
            if not total:
                continue
            top = scoped.values('publisher__name').annotate(count=Count('id')).order_by('-count').first()
            latest = scoped.order_by('-publication_date').values_list('publication_date', flat=True).first()
            comparison.append({
                'country': item.name,
                'code': item.code,
                'flag': item.flag,
                'total': total,
                'positive': scoped.filter(sentiment='Positive').count(),
                'negative': scoped.filter(sentiment='Negative').count(),
                'neutral': scoped.filter(sentiment='Neutral').count(),
                'top_publisher': top['publisher__name'] if top else '—',
                'latest': latest,
            })

        buckets = defaultdict(lambda: {'articles': 0, 'Positive': 0, 'Negative': 0, 'Neutral': 0})
        for row in qs.values('publication_date', 'sentiment').annotate(count=Count('id')):
            day = row['publication_date']
            if not day:
                continue
            buckets[day]['articles'] += row['count']
            if row['sentiment'] in buckets[day]:
                buckets[day][row['sentiment']] += row['count']

        series = []
        cursor = start
        while cursor <= end:
            point = buckets[cursor]
            series.append({
                'date': cursor.isoformat(),
                'label': cursor.strftime('%d %b'),
                'articles': point['articles'],
                'positive': point['Positive'],
                'negative': point['Negative'],
                'neutral': point['Neutral'],
            })
            cursor += timedelta(days=1)

        volumes = [point['articles'] for point in series]
        for index, point in enumerate(series):
            window = volumes[max(0, index - 2): index + 1]
            point['ma3'] = round(sum(window) / len(window), 2) if window else 0

        prev_start = start - timedelta(days=days)
        previous_total = Content.objects.filter(publication_date__gte=prev_start, publication_date__lt=start)
        if country:
            previous_total = previous_total.filter(country=country)
        previous_total = previous_total.count()
        current_total = qs.count()
        peak = max(series, key=lambda point: point['articles']) if series else None
        first_half = sum(volumes[: max(len(volumes) // 2, 1)])
        second_half = sum(volumes[len(volumes) // 2:])
        trend = 'rising' if second_half > first_half else 'easing' if second_half < first_half else 'stable'

        return Response({
            'scope': country.name if country else 'Global',
            'period_days': days,
            'total': current_total,
            'previous_total': previous_total,
            'delta': current_total - previous_total,
            'peak_day': peak['date'] if peak and peak['articles'] else None,
            'peak_count': peak['articles'] if peak else 0,
            'trend': trend,
            'by_country': by_country,
            'sentiment': sentiments,
            'comparison': comparison,
            'series': series,
        })


class FetchJobListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return paginate(FetchJob.objects.all(), request, FetchJobSerializer)

    def post(self, request):
        if has_active_fetch_job():
            return Response({'error': 'A fetch job is already running.'}, status=400)
        today = timezone.now().date()
        payload = {
            'keywords': request.data.get('keywords') or '',
            'categories': ','.join(request.data.get('categories') or []),
            'start_date': request.data.get('start_date') or today - timedelta(days=3),
            'end_date': request.data.get('end_date') or today,
            'publisher_ids': ','.join(str(pid) for pid in (request.data.get('publisher_ids') or [])),
            'include_newsnow': bool(request.data.get('include_newsnow', True)),
            'include_google': bool(request.data.get('include_google', True)),
        }
        job = FetchJob.objects.create(created_by=request.user, **payload)
        start_fetch_job(job.id)
        return Response(FetchJobSerializer(job).data, status=status.HTTP_201_CREATED)


class FetchJobDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        return Response(FetchJobSerializer(get_object_or_404(FetchJob, pk=pk)).data)


class ExportWordView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        articles = Content.objects.select_related('country', 'publisher')[:400]
        try:
            from docx import Document
        except ImportError:
            lines = ['NarrativeShield media corpus\n']
            for article in articles:
                lines.append(f'{article.publication_date} | {article.country.name} | {article.publisher.name}')
                lines.append(article.title)
                lines.append(article.summary)
                lines.append(article.url)
                lines.append('')
            response = HttpResponse('\n'.join(lines), content_type='text/plain; charset=utf-8')
            response['Content-Disposition'] = 'attachment; filename="narrativeshield-media.txt"'
            return response

        document = Document()
        document.add_heading('NarrativeShield · Media corpus', 0)
        document.add_paragraph(f'Exported {timezone.now().isoformat()} · {articles.count()} articles')
        for article in articles:
            document.add_heading(article.title, level=2)
            document.add_paragraph(
                f'{article.publication_date} · {article.country.name} · {article.publisher.name} · {article.sentiment or "—"} · {article.category or "—"}'
            )
            document.add_paragraph(article.summary or article.body[:400])
            document.add_paragraph(article.url)
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        response = HttpResponse(buffer.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = 'attachment; filename="narrativeshield-media.docx"'
        return response
