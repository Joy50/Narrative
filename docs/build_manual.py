from pathlib import Path
import re
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Preformatted, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parent
source = (ROOT / 'NarrativeShield_User_Manual.md').read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.78)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.10
doc.styles['Title'].font.size = Pt(28)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(12)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('NarrativeShield User Manual  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(8)

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='ManualBody', fontName='Helvetica', fontSize=10, leading=13.2, spaceAfter=6))
styles.add(ParagraphStyle(name='ManualTitle', fontName='Helvetica-Bold', fontSize=27, leading=32, spaceAfter=20, textColor=colors.black))
styles.add(ParagraphStyle(name='ManualH1', fontName='Helvetica-Bold', fontSize=19, leading=24, spaceAfter=15, keepWithNext=True))
styles.add(ParagraphStyle(name='ManualH2', fontName='Helvetica-Bold', fontSize=11.5, leading=14, spaceBefore=8, spaceAfter=5, keepWithNext=True))
styles.add(ParagraphStyle(name='ManualList', parent=styles['ManualBody'], leftIndent=14, firstLineIndent=-10, spaceAfter=5))
styles.add(ParagraphStyle(name='ManualCode', fontName='Courier', fontSize=8.5, leading=12, spaceBefore=4, spaceAfter=9))
story = []
in_code = False
code = []
toc = False
for line in source.splitlines():
    if line.startswith('```'):
        if not in_code:
            in_code = True; code = []
        else:
            p = doc.add_paragraph('\n'.join(code))
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
            story.append(Preformatted('\n'.join(code), styles['ManualCode']))
            in_code = False
        continue
    if in_code:
        code.append(line); continue
    if not line.strip(): continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
        story.append(Spacer(1, 30))
        story.append(Paragraph(escape(line[2:]), styles['ManualTitle']))
    elif line.startswith('## '):
        heading = line[3:]
        doc.add_page_break()
        doc.add_paragraph(heading, 'Heading 1')
        story.append(PageBreak())
        story.append(Paragraph(escape(heading), styles['ManualH1']))
        toc = heading == 'Contents'
    elif line.startswith('### '):
        if line[4:] == 'Account administration and review access':
            doc.add_page_break()
            story.append(PageBreak())
        doc.add_paragraph(line[4:], 'Heading 2')
        story.append(Paragraph(escape(line[4:]), styles['ManualH2']))
    else:
        listing = line.startswith('- ') or re.match(r'^\d+\. ', line)
        text = line.replace('- ', '\u2022 ', 1) if line.startswith('- ') else line
        p = doc.add_paragraph(text)
        if listing:
            p.paragraph_format.left_indent = Inches(.17)
            p.paragraph_format.first_line_indent = Inches(-.13)
        story.append(Paragraph(escape(text), styles['ManualList'] if listing else styles['ManualBody']))
doc.core_properties.title = 'NarrativeShield User Manual'
doc.core_properties.subject = 'User and administrator guide'
doc.core_properties.author = 'NarrativeShield'
doc.save(ROOT / 'NarrativeShield_User_Manual.docx')

def page(canvas, document):
    canvas.setTitle('NarrativeShield User Manual')
    canvas.setAuthor('NarrativeShield')
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#58656A'))
    canvas.drawString(56, 28, 'NarrativeShield  |  User and administrator guide')
    canvas.drawRightString(A4[0]-56, 28, str(document.page))

pdf = SimpleDocTemplate(str(ROOT / 'NarrativeShield_User_Manual.pdf'), pagesize=A4,
    rightMargin=56, leftMargin=56, topMargin=48, bottomMargin=48)
pdf.build(story, onFirstPage=page, onLaterPages=page)
import pymupdf
pdf_path = ROOT / 'NarrativeShield_User_Manual.pdf'
rendered = pymupdf.open(pdf_path)
chapters = [line[3:] for line in source.splitlines() if re.match(r'^## \d+ ', line)]
outline = []
for chapter in chapters:
    for i, pg in enumerate(rendered):
        if chapter in ' '.join(pg.get_text().split()):
            outline.append([1, chapter, i+1]); break
rendered.set_toc(outline)
contents = rendered[1]
for _, chapter, number in outline:
    entry = re.sub(r'^(\d+) ', r'\1. ', chapter)
    for rect in contents.search_for(entry):
        contents.insert_text((520, rect.y1-2), str(number), fontsize=10, fontname='helv')
        contents.insert_link({'kind': pymupdf.LINK_GOTO, 'from': pymupdf.Rect(rect.x0, rect.y0, 540, rect.y1), 'page': number-1})
temporary = ROOT / 'manual_numbered.pdf'
rendered.save(temporary)
rendered.close()
temporary.replace(pdf_path)
print('Created Word and PDF manuals')
